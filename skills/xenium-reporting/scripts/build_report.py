#!/usr/bin/env python3
"""Assemble a native Word scientific report from source-bound content blocks.

The optional renderer must implement the render_docx.py command interface.
Tables are formatted, never statistically recomputed. Paths in content are relative
to the content file. A fresh output directory keeps reviewed builds immutable.
"""
import argparse
import csv
import json
import os
import platform
import re
import subprocess
import shlex
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

import docx
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image


def require(condition, message):
    if not condition:
        raise ValueError(message)


def resolve_source(base, value):
    require(isinstance(value, str) and value.strip(), 'Source path must be a nonempty string')
    p = (base / value).resolve()
    require(p.is_file(), f'Missing source: {p}')
    return p


def format_value(value, kind):
    if kind == 'text':
        return value
    require(kind in ('integer', 'fixed2', 'fixed3'), f'Unknown numeric format: {kind}')
    try:
        n = Decimal(value)
    except InvalidOperation as e:
        raise ValueError(f'Invalid numeric value: {value!r}') from e
    require(n.is_finite(), f'Nonfinite numeric value: {value!r}; declare missingness explicitly in the source')
    if kind == 'integer':
        require(n == n.to_integral_value(), f'Noninteger count: {value!r}')
        return f'{int(n):,}'
    return format(n, f'.{kind[-1]}f')


def prepare(content_path, style):
    content = json.loads(content_path.read_text())
    require(isinstance(content.get('title'), str) and content['title'].strip(), 'Report title is required')
    require(isinstance(content.get('blocks'), list) and content['blocks'], 'Content blocks are required')
    cfg = style['report']
    text_width = cfg['page']['width_mm'] - 2 * cfg['page']['side_margin_mm']
    blocks = []
    for i, block in enumerate(content['blocks']):
        b = dict(block)
        kind = b.get('type')
        require(kind in ('heading', 'paragraph', 'table', 'figure', 'page_break'), f'Unknown block type at {i}: {kind}')
        if kind in ('heading', 'paragraph'):
            require(isinstance(b.get('text'), str) and b['text'].strip(), f'Empty text block {i}')
        if kind == 'heading':
            require(b.get('level', 1) in (1, 2), 'Only heading levels 1 and 2 are supported')
        if kind == 'table':
            source = resolve_source(content_path.parent, b['source'])
            require(b.get('columns'), f'Table {i} needs columns')
            delimiter = b.get('delimiter', '\t')
            require(len(delimiter) == 1, 'Delimiter must be one character')
            with source.open(newline='') as f:
                reader = csv.DictReader(f, delimiter=delimiter)
                fields = reader.fieldnames or []
                require(len(fields) == len(set(fields)), f'Duplicate source columns: {source}')
                for col in b['columns']:
                    require(col['key'] in fields, f'Missing column {col["key"]!r} in {source}')
                    require(float(col.get('width', 1)) > 0, 'Column widths must be positive')
                rows = list(reader)
            require(rows, f'Empty table: {source}')
            require(all(None not in row and all(v is not None for v in row.values()) for row in rows), f'Malformed row: {source}')
            if 'expected_rows' in b:
                require(len(rows) == b['expected_rows'], f'Unexpected row count: {source}')
            b['_rows'] = [[format_value(row[col['key']], col.get('format', 'text')) for col in b['columns']] for row in rows]
            b['_source'] = str(source)
        if kind == 'figure':
            source = resolve_source(content_path.parent, b['source'])
            require(b.get('caption'), f'Figure {i} needs a separate caption')
            width = float(b.get('width_mm', cfg['figure']['width_mm']))
            require(0 < width <= text_width + .01, f'Figure width {width} exceeds text width {text_width}')
            with Image.open(source) as im:
                require(im.format == 'PNG', 'Use a PNG report image; retain vector masters separately')
                pixels = list(im.size)
            dpi = pixels[0] / (width / 25.4)
            require(dpi >= style['dpi'] - 1, f'Figure has only {dpi:.1f} effective dpi at {width} mm')
            b.update(_source=str(source), _pixels=pixels, _effective_dpi=dpi, width_mm=width)
        blocks.append(b)
    provenance = []
    for value in content.get('provenance', []):
        # Metadata records may include a description alongside the source path.
        p = value.get('source', value.get('path')) if isinstance(value, dict) else value
        provenance.append(str(resolve_source(content_path.parent, p)))
    return content, blocks, provenance


def set_font(style, name, size, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.complex_script = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn(f'w:{key}'), None)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn(f'w:{key}'), name)


def named_style(doc, name):
    return doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def setup_styles(doc, style):
    cfg = style['report']
    typo = cfg['typography']
    family = cfg['font_family']
    definitions = {
        'Normal': ('body_pt', False), 'Title': ('title_pt', True),
        'Heading 1': ('heading1_pt', True), 'Heading 2': ('heading2_pt', True),
        'Caption': ('caption_pt', False), 'Table Text': ('table_pt', False),
        'Table Note': ('caption_pt', False), 'Header': ('footer_pt', False),
        'Footer': ('footer_pt', False),
    }
    for name, (key, bold) in definitions.items():
        s = named_style(doc, name)
        set_font(s, family, typo[key], bold)
        pf = s.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(typo['paragraph_after_pt'] if name == 'Normal' else 0)
        pf.line_spacing = typo['line_spacing'] if name == 'Normal' else 1.1
        pf.widow_control = True
        pf.keep_with_next = name.startswith('Heading') or name == 'Title'
        pf.keep_together = name in ('Caption', 'Table Note')
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for border in list(s.element.xpath('./w:pPr/w:pBdr')):
            border.getparent().remove(border)
    doc.styles['Title'].paragraph_format.space_after = Pt(10)
    for level, before, after in ((1, 14, 6), (2, 10, 4)):
        doc.styles[f'Heading {level}'].paragraph_format.space_before = Pt(before)
        doc.styles[f'Heading {level}'].paragraph_format.space_after = Pt(after)
    doc.styles['Caption'].paragraph_format.space_after = Pt(6)


def add_rich_lead(paragraph, text):
    # Only the first caption/title sentence is emphasized; no Markdown parser.
    match = re.match(r'((?:Figure|Table) \d+\. [^.]+\.)(.*)', text, re.S)
    if match:
        paragraph.add_run(match.group(1)).bold = True
        paragraph.add_run(match.group(2))
    else:
        paragraph.add_run(text)


def cell_border(cell, edge, value='nil', size='4'):
    pr = cell._tc.get_or_add_tcPr()
    borders = pr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        pr.append(borders)
    e = OxmlElement(f'w:{edge}')
    e.set(qn('w:val'), value)
    e.set(qn('w:sz'), size)
    e.set(qn('w:color'), '000000')
    borders.append(e)


def add_table(doc, block, width_mm):
    title = doc.add_paragraph(style='Caption')
    add_rich_lead(title, block['title'])
    title.paragraph_format.keep_with_next = True
    title.paragraph_format.space_before = Pt(8)
    table = doc.add_table(rows=1, cols=len(block['columns']))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    columns = block['columns']
    weights = sum(float(c.get('width', 1)) for c in columns)
    widths = [width_mm * float(c.get('width', 1)) / weights for c in columns]
    for col, width in zip(table.columns, widths):
        col.width = Mm(width)
    header = table.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    data = [[c['label'] for c in columns], *block['_rows']]
    for ri, values in enumerate(data):
        row = header if ri == 0 else table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Mm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.style = doc.styles['Table Text']
            p.paragraph_format.keep_with_next = ri == 0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if columns[j].get('format', 'text') == 'text' else WD_ALIGN_PARAGRAPH.RIGHT
            p.add_run(value).bold = ri == 0
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for tag, amount in (('top', 85), ('bottom', 85), ('left', 65), ('right', 65)):
                el = OxmlElement(f'w:{tag}')
                el.set(qn('w:w'), str(amount))
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            pr.append(margins)
            for edge in ('top', 'bottom', 'left', 'right'):
                drawn = (edge == 'top' and ri == 0) or (edge == 'bottom' and ri in (0, len(data) - 1))
                cell_border(cell, edge, 'single' if drawn else 'nil')
    if block.get('note'):
        note = doc.add_paragraph(block['note'], style='Table Note')
        note.paragraph_format.space_before = Pt(5)
        note.paragraph_format.space_after = Pt(6)


def build_document(content, blocks, style, output):
    doc = Document()
    setup_styles(doc, style)
    page = style['report']['page']
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(page['width_mm']), Mm(page['height_mm'])
    section.left_margin = section.right_margin = Mm(page['side_margin_mm'])
    section.top_margin = section.bottom_margin = Mm(page['vertical_margin_mm'])
    section.header_distance = section.footer_distance = Mm(9)
    doc.core_properties.title = content['title']
    doc.core_properties.subject = content.get('subject', 'Exploratory scientific analysis')
    doc.core_properties.author = content.get('author', '')
    doc.core_properties.last_modified_by = content.get('author', '')
    doc.add_paragraph(content['title'], style='Title')
    footer = section.footer.paragraphs[0]
    footer.style = doc.styles['Footer']
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(content.get('short_title', '') + '   |   ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    width = page['width_mm'] - 2 * page['side_margin_mm']
    for block in blocks:
        kind = block['type']
        if kind == 'paragraph':
            doc.add_paragraph(block['text'])
        elif kind == 'heading':
            doc.add_paragraph(block['text'], style=f'Heading {block.get("level", 1)}')
        elif kind == 'table':
            add_table(doc, block, width)
        elif kind == 'figure':
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1
            shape = p.add_run().add_picture(block['_source'], width=Mm(block['width_mm']))
            shape._inline.docPr.set('descr', block.get('alt', block['caption']))
            caption = doc.add_paragraph(style='Caption')
            add_rich_lead(caption, block['caption'])
        elif kind == 'page_break':
            doc.add_page_break()
    doc.save(output)


def pdf_details(pdf):
    from pypdf import PdfReader
    reader = PdfReader(pdf)
    fonts = {}
    sizes = []
    text = []
    for page in reader.pages:
        sizes.append([float(page.mediabox.width), float(page.mediabox.height)])
        text.append(page.extract_text() or '')
        resources = page.get('/Resources', {})
        resources = resources.get_object() if hasattr(resources, 'get_object') else resources
        for ref in resources.get('/Font', {}).values():
            font = ref.get_object()
            name = str(font.get('/BaseFont', ''))
            targets = font.get('/DescendantFonts', [font])
            for target in targets:
                target = target.get_object() if hasattr(target, 'get_object') else target
                descriptor = target.get('/FontDescriptor', {})
                descriptor = descriptor.get_object() if hasattr(descriptor, 'get_object') else descriptor
                embedded = any(k in descriptor for k in ('/FontFile', '/FontFile2', '/FontFile3'))
                fonts[name] = {'name': name, 'embedded': embedded}
    return {'page_count': len(reader.pages), 'page_sizes_pt': sizes, 'fonts': list(fonts.values())}, '\n\n'.join(text)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--content', type=Path, required=True)
    ap.add_argument('--style', type=Path, default=Path(__file__).resolve().parents[2] / 'xenium-figures/assets/style.json')
    ap.add_argument('--output-dir', type=Path, required=True, help='New directory; existing outputs are never overwritten')
    ap.add_argument('--fontconfig', type=Path, help='Existing fontconfig file for this conversion process; no global font changes')
    group = ap.add_mutually_exclusive_group(required=True)
    group.add_argument('--renderer', type=Path, help='Environment-provided render_docx.py; uses this Python runtime')
    group.add_argument('--docx-only', action='store_true', help='Structure/debug mode; not a reviewed report delivery')
    args = ap.parse_args()
    content_path, style_path, out = args.content.resolve(), args.style.resolve(), args.output_dir.resolve()
    require(content_path.is_file(), f'Missing content: {content_path}')
    require(not out.exists(), f'Output directory already exists: {out}; use a new build directory')
    style = json.loads(style_path.read_text())
    content, blocks, provenance = prepare(content_path, style)
    if args.renderer:
        require(args.renderer.is_file(), f'Missing renderer: {args.renderer}')
    if args.fontconfig:
        require(args.fontconfig.is_file(), f'Missing fontconfig: {args.fontconfig}')
    out.mkdir(parents=True)
    output = out / 'report.docx'
    manifest = {
        'status': 'building', 'content': str(content_path), 'style_source': str(style_path),
        'effective_style': style, 'provenance': provenance, 'docx': str(output),
        'versions': {'python': platform.python_version(), 'python_docx': docx.__version__},
        'command': [sys.executable, str(Path(__file__).resolve()), *sys.argv[1:]],
        'tables': [{'source': b['_source'], 'columns': b['columns'], 'rows': b['_rows']} for b in blocks if b['type'] == 'table'],
        'figures': [{'source': b['_source'], 'width_mm': b['width_mm'], 'pixels': b['_pixels'], 'effective_dpi': b['_effective_dpi']} for b in blocks if b['type'] == 'figure'],
        'additional_review': 'not_requested',
    }
    manifest_path = out / 'manifest.json'
    try:
        build_document(content, blocks, style, output)
        manifest['status'] = 'docx_only'
        if args.renderer:
            render_dir = out / 'render'
            cmd = [sys.executable, str(args.renderer.resolve()), str(output), '--output_dir', str(render_dir), '--emit_pdf', '--verbose', '--dpi', '144']
            manifest['conversion'] = {'input_docx': str(output), 'command': cmd, 'log': str(out / 'conversion.log')}
            env = os.environ.copy()
            if args.fontconfig:
                env['FONTCONFIG_FILE'] = str(args.fontconfig.resolve())
            manifest['conversion']['fontconfig'] = env.get('FONTCONFIG_FILE')
            proc = subprocess.run(cmd, env=env, text=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
            (out / 'conversion.log').write_text(proc.stdout)
            require(proc.returncode == 0, f'Render failed; see {out / "conversion.log"}')
            require(not re.search(r'Repaired OOXML|DOCX->ODT|ODT->PDF|--convert-to\s+odt', proc.stdout), 'Renderer used repair or ODT fallback; direct conversion not established')
            pdf = render_dir / 'report.pdf'
            require(pdf.is_file(), 'Renderer did not preserve the PDF used for page previews')
            details, pdf_text = pdf_details(pdf)
            for line in proc.stdout.splitlines():
                if line.startswith('[render_docx] $ '):
                    converter = shlex.split(line[len('[render_docx] $ '):])[0]
                    version = subprocess.run([converter, '--version'], env=env, text=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=30)
                    manifest['versions']['converter'] = version.stdout.strip()
                    break
            (out / 'pdf-text.txt').write_text(pdf_text)
            pages = sorted(render_dir.glob('page-*.png'), key=lambda p: int(p.stem.split('-')[-1]))
            require(len(pages) == details['page_count'], 'Page preview count differs from PDF page count')
            manifest['conversion'].update({'pdf': str(pdf), 'pages': list(map(str, pages)), **details})
            fonts = details['fonts']
            family = style['report']['font_family'].replace(' ', '')
            face = re.compile(re.escape(family) + r'(?:(?:-|,)?(?:BoldItalic|BoldOblique|Regular|Bold|Italic|Oblique))?(?:MT)?')
            require(fonts and all(face.fullmatch(f['name'].lstrip('/').split('+')[-1].replace(' ', '')) and f['embedded'] for f in fonts), f'PDF must embed actual {family}: {fonts}')
            manifest['conversion'].update({'route': 'direct_docx_to_pdf', 'pdf': str(pdf), 'pages': list(map(str, pages)), **details})
            manifest['status'] = 'built'
    except Exception as exc:
        manifest['status'] = 'failed'
        manifest['error'] = str(exc)
        raise
    finally:
        manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False, allow_nan=False) + '\n')
    print(json.dumps({'status': manifest['status'], 'docx': str(output), 'manifest': str(manifest_path)}))


if __name__ == '__main__':
    try:
        main()
    except (ValueError, KeyError, OSError, json.JSONDecodeError) as exc:
        print(f'ERROR: {exc}', file=sys.stderr)
        sys.exit(2)
