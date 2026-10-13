"""Write a shareable copy of a built deck: speaker notes removed, file properties set.

Run with the existing python-pptx environment: python share_copy.py BUILT.pptx SHARED.pptx \
  [--author NAME] [--title TITLE]

The built deck keeps its notes; the copy goes to a new path and the script refuses an
existing destination. Exit status 1 means no shareable copy was written.
"""
import argparse
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("built", type=Path)
    ap.add_argument("shared", type=Path)
    ap.add_argument("--author", help="author and last-modified-by; builders often leave a tool name here")
    ap.add_argument("--title", help="document title property")
    a = ap.parse_args()
    if a.shared.exists():
        sys.exit(f"{a.shared} exists; choose a new name")
    prs = Presentation(a.built)
    for slide in prs.slides:
        for rid, rel in list(slide.part.rels.items()):
            if rel.reltype == RT.NOTES_SLIDE:
                slide.part.drop_rel(rid)
    props = prs.core_properties
    if a.author:
        props.author = props.last_modified_by = a.author
    if a.title:
        props.title = a.title
    with a.shared.open("xb") as output:
        prs.save(output)
    with zipfile.ZipFile(a.shared) as package:
        left = [n for n in package.namelist() if n.startswith("ppt/notesSlides/")]
    if left or any(s.has_notes_slide for s in Presentation(a.shared).slides):
        a.shared.unlink()
        sys.exit("notes remain in the copy; nothing written")
    print(f"{a.shared}: {len(prs.slides)} slides, no notes; author {props.author!r}, title {props.title!r}")


if __name__ == "__main__":
    main()
