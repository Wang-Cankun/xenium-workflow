"""Exercise the sharing CLI on a disposable deck with notes and native content."""
import base64
import io
from pathlib import Path
import subprocess
import sys
import tempfile
import zipfile

from pptx import Presentation
from pptx.util import Inches

SCRIPT = Path(__file__).resolve().parents[1] / 'skills/xenium-reporting/scripts/share_copy.py'


def main():
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        built, shared = root / 'built.pptx', root / 'shared.pptx'
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = 'Coverage by region'
        slide.notes_slide.notes_text_frame.text = 'Private speaking cues'
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        table.cell(0, 0).text, table.cell(1, 0).text = 'Region', 'Region A'
        table.cell(1, 1).text = '12'
        png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=')
        slide.shapes.add_picture(io.BytesIO(png), Inches(1), Inches(4))
        deck.save(built)
        original = built.read_bytes()
        command = [sys.executable, str(SCRIPT), str(built), str(shared),
                   '--author', 'Presenter', '--title', 'Progress briefing']
        result = subprocess.run(command, capture_output=True, text=True)
        assert result.returncode == 0, result.stderr
        assert built.read_bytes() == original
        result_deck = Presentation(shared)
        assert len(result_deck.slides) == 1
        assert not result_deck.slides[0].has_notes_slide
        assert result_deck.core_properties.author == 'Presenter'
        assert result_deck.core_properties.last_modified_by == 'Presenter'
        assert result_deck.core_properties.title == 'Progress briefing'
        with zipfile.ZipFile(built) as source, zipfile.ZipFile(shared) as output:
            assert not any(n.startswith('ppt/notesSlides/') for n in output.namelist())
            for name in source.namelist():
                if name.startswith(('ppt/slides/slide', 'ppt/media/', 'ppt/slideLayouts/', 'ppt/slideMasters/')):
                    assert output.read(name) == source.read(name), name
        delivered = shared.read_bytes()
        refused = subprocess.run(command, capture_output=True, text=True)
        assert refused.returncode != 0 and 'exists' in refused.stderr
        assert shared.read_bytes() == delivered
        assert built.read_bytes() == original
    print('Sharing CLI: notes removed, content and original preserved, metadata set, overwrite refused')


if __name__ == '__main__':
    main()
